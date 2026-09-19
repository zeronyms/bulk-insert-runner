import csv
import json
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def flatten_dict(d: dict[str, Any], parent_key: str = "", sep: str = ".") -> dict[str, Any]:
    """Flatten nested dict into dot notation.

    Example: {'user': {'name': 'John'}} -> {'user.name': 'John'}
    """
    items: list[tuple[str, Any]] = []
    for k, v in d.items():
        new_key = f"{parent_key}{sep}{k}" if parent_key else str(k)
        if isinstance(v, dict) and v:
            items.extend(flatten_dict(v, new_key, sep=sep).items())
        else:
            items.append((new_key, v))
    return dict(items)


def unflatten_dict(d: dict[str, Any], sep: str = ".") -> dict[str, Any]:
    """Unflatten dot-notation dict back into nested dict."""
    result: dict[str, Any] = {}
    for key, value in d.items():
        parts = key.split(sep)
        curr = result
        for part in parts[:-1]:
            if part not in curr or not isinstance(curr[part], dict):
                curr[part] = {}
            curr = curr[part]
        curr[parts[-1]] = value
    return result


def serialize_cell_value(val: Any) -> Any:
    """Serialize lists and dicts into JSON strings for tabular cells."""
    if isinstance(val, (list, dict)):
        return json.dumps(val, ensure_ascii=False)
    return val


def cast_cell_value(val: Any, template_val: Any = None) -> Any:
    """Restore cell value into appropriate native Python type based on template schema."""
    if val is None or (isinstance(val, str) and val.strip() == ""):
        if template_val is None:
            return None
        if isinstance(template_val, str) and template_val == "":
            return ""
        return None

    # Boolean handling
    if isinstance(template_val, bool):
        if isinstance(val, bool):
            return val
        s = str(val).strip().lower()
        if s in ("true", "1", "yes", "ya", "t", "y"):
            return True
        if s in ("false", "0", "no", "tidak", "f", "n"):
            return False
        return bool(val)

    # Integer handling (avoid bool as bool is a subclass of int)
    if isinstance(template_val, int) and not isinstance(template_val, bool):
        try:
            return int(round(float(str(val).strip())))
        except (ValueError, TypeError):
            return val

    # Float handling
    if isinstance(template_val, float):
        try:
            return float(str(val).strip())
        except (ValueError, TypeError):
            return val

    # List or Dict handling
    if isinstance(template_val, (list, dict)):
        if isinstance(val, (list, dict)):
            return val
        if isinstance(val, str):
            val_str = val.strip()
            if (val_str.startswith("[") and val_str.endswith("]")) or (
                val_str.startswith("{") and val_str.endswith("}")
            ):
                try:
                    return json.loads(val_str)
                except json.JSONDecodeError:
                    pass
        return template_val

    # Generic heuristics if no template reference was provided
    if isinstance(val, str):
        val_str = val.strip()
        if (val_str.startswith("[") and val_str.endswith("]")) or (
            val_str.startswith("{") and val_str.endswith("}")
        ):
            try:
                return json.loads(val_str)
            except json.JSONDecodeError:
                pass

        if val_str.lower() in ("true", "false"):
            return val_str.lower() == "true"

    return val


def _extract_headers_and_flatten(items: list[Any]) -> tuple[list[str], list[dict[str, Any]]]:
    """Helper to flatten items and extract unique headers in order."""
    flattened_items: list[dict[str, Any]] = []
    headers: list[str] = []
    header_set: set[str] = set()

    for item in items:
        if isinstance(item, dict):
            flat = flatten_dict(item)
            flattened_items.append(flat)
            for k in flat:
                if k not in header_set:
                    header_set.add(k)
                    headers.append(k)
        else:
            headers = ["value"]
            flattened_items.append({"value": item})

    return headers, flattened_items


def export_excel_template(filepath: Path | str, items: list[Any]) -> bool:
    """Export items into an Excel (.xlsx) template file with styled headers."""
    headers, flattened_items = _extract_headers_and_flatten(items)
    if not headers:
        return False

    wb = openpyxl.Workbook()
    ws = wb.active
    if ws is None:
        return False

    ws.title = "Bulk Insert"
    ws.append(headers)

    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for flat_item in flattened_items:
        row_vals = [serialize_cell_value(flat_item.get(h)) for h in headers]
        ws.append(row_vals)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = thin_border

    ws.freeze_panes = "A2"

    for col in ws.columns:
        col_idx = col[0].column
        if isinstance(col_idx, int):
            col_letter = get_column_letter(col_idx)
            max_len = max((len(str(cell.value or "")) for cell in col), default=0)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)

    wb.save(filepath)
    return True


def export_csv_template(filepath: Path | str, items: list[Any]) -> bool:
    """Export items into a CSV template file with UTF-8 BOM encoding."""
    headers, flattened_items = _extract_headers_and_flatten(items)
    if not headers:
        return False

    with open(filepath, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for flat in flattened_items:
            row_vals = [serialize_cell_value(flat.get(h)) for h in headers]
            writer.writerow(row_vals)

    return True


def export_json_template(filepath: Path | str, items: list[Any]) -> None:
    """Export items into a JSON template file."""
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(items, f, indent=2, ensure_ascii=False)


def load_excel_data(filepath: Path | str, sample_payload: Any = None) -> list[dict[str, Any]]:
    """Load and parse data rows from an Excel (.xlsx) file."""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active
    if ws is None:
        return []

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    if not any(headers):
        return []

    sample_flat: dict[str, Any] = {}
    if sample_payload:
        sample_item = (
            sample_payload[0]
            if isinstance(sample_payload, list) and sample_payload
            else sample_payload
        )
        if isinstance(sample_item, dict):
            sample_flat = flatten_dict(sample_item)

    data_items: list[dict[str, Any]] = []
    for row in rows[1:]:
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        flat_row: dict[str, Any] = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            cell_val = row[idx] if idx < len(row) else None
            tmpl_val = sample_flat.get(header)
            flat_row[header] = cast_cell_value(cell_val, tmpl_val)

        data_items.append(unflatten_dict(flat_row))

    return data_items


def load_csv_data(filepath: Path | str, sample_payload: Any = None) -> list[dict[str, Any]]:
    """Load and parse data rows from a CSV file (auto-detecting delimiter)."""
    sample_flat: dict[str, Any] = {}
    if sample_payload:
        sample_item = (
            sample_payload[0]
            if isinstance(sample_payload, list) and sample_payload
            else sample_payload
        )
        if isinstance(sample_item, dict):
            sample_flat = flatten_dict(sample_item)

    with open(filepath, encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            delimiter = dialect.delimiter
        except Exception:
            delimiter = ","

        reader = csv.DictReader(f, delimiter=delimiter)
        data_items: list[dict[str, Any]] = []
        for row in reader:
            if not row or all(v is None or str(v).strip() == "" for v in row.values()):
                continue
            flat_row: dict[str, Any] = {}
            for k, v in row.items():
                if k is None or not k.strip():
                    continue
                k_clean = k.strip()
                tmpl_val = sample_flat.get(k_clean)
                flat_row[k_clean] = cast_cell_value(v, tmpl_val)
            data_items.append(unflatten_dict(flat_row))

    return data_items


def load_json_data(filepath: Path | str) -> list[dict[str, Any]]:
    """Load data rows from a JSON file."""
    with open(filepath, encoding="utf-8") as f:
        data = json.load(f)
    if isinstance(data, list):
        return data
    if isinstance(data, dict):
        return [data]
    raise TypeError("JSON file format must be an array [] or an object {}")


def load_data_file(filepath: Path | str, sample_payload: Any = None) -> list[dict[str, Any]]:
    """Universal loader that detects file extension (.xlsx, .csv, .json) and parses records."""
    path = Path(filepath)
    suffix = path.suffix.lower()

    if suffix == ".xlsx":
        return load_excel_data(path, sample_payload)
    if suffix == ".csv":
        return load_csv_data(path, sample_payload)
    if suffix == ".json":
        return load_json_data(path)

    raise ValueError(f"Unsupported file format: {suffix} (only .xlsx, .csv, .json supported)")


def format_template_str(template_str: str, values: dict[str, Any]) -> str:
    """Safely replace {key} in template_str using values dictionary without KeyError."""
    res = template_str
    for k, v in values.items():
        placeholder = f"{{{k}}}"
        if placeholder in res:
            res = res.replace(placeholder, str(v))
    return res


def group_items_for_request(
    items: list[dict[str, Any]],
    group_by: str | None,
    array_key: str,
) -> list[tuple[dict[str, Any], dict[str, Any]]]:
    """Group flat items from Excel/CSV into batch requests.

    Returns a list of tuples: (context_vars, request_payload)
    - context_vars: e.g. {"learningPlanSerial": "LRPN-RZ48G8E4"}
    - request_payload: e.g. {"courses": [course1, course2, ...]}
    """
    if not group_by:
        return [({}, {array_key: items})]

    groups: dict[str, list[dict[str, Any]]] = {}
    for item in items:
        raw_key = item.get(group_by)
        if raw_key is None or str(raw_key).strip() == "":
            continue
        key_str = str(raw_key).strip()
        if key_str not in groups:
            groups[key_str] = []
        clean_item = {k: v for k, v in item.items() if k != group_by}
        groups[key_str].append(clean_item)

    result: list[tuple[dict[str, Any], dict[str, Any]]] = []
    for key_str, group_rows in groups.items():
        context_vars = {group_by: key_str}
        request_payload = {array_key: group_rows}
        result.append((context_vars, request_payload))

    return result
